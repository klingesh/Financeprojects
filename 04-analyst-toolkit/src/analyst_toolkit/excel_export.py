"""Table export to CSV and, optionally, formatted Excel.

CSV works with no dependencies. Excel export requires ``openpyxl``; if it is
absent the caller gets a clear instruction rather than an ImportError
traceback.
"""

import csv
from typing import Any, List, Optional, Sequence

# Fields rendered as multiples (1 decimal, x suffix) vs. percentages.
MULTIPLE_LABELS = {
    "EV/Revenue",
    "EV/EBITDA",
    "EV/EBIT",
    "P/E",
    "P/TBV",
    "Net debt/EBITDA",
}
PERCENT_LABELS = {"EBITDA margin", "ROTE", "P/AUM"}

# Rendered as a plain number to one decimal place — basis points and per-unit
# currency figures are neither multiples nor percentages.
DECIMAL_LABELS = {"Blended yield (bps)", "EV per account", "Revenue per account"}


def format_value(label: str, value: Any) -> str:
    """Render a cell for human-readable output."""
    if value is None:
        return "n/a"
    if not isinstance(value, (int, float)):
        return str(value)

    if label in PERCENT_LABELS:
        return "{0:.1f}%".format(value * 100)
    if label in MULTIPLE_LABELS:
        return "{0:.1f}x".format(value)
    if label in DECIMAL_LABELS:
        return "{0:,.1f}".format(value)
    if abs(value) >= 1000:
        return "{0:,.0f}".format(value)
    return "{0:,.2f}".format(value)


def write_csv(table: Sequence[Sequence[Any]], path: str, raw: bool = False) -> str:
    """Write a table to CSV.

    ``raw=True`` writes unformatted numbers, which is what you want if the
    output feeds another tool. Default writes formatted display strings.
    """
    header = list(table[0]) if table else []
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for index, row in enumerate(table):
            if index == 0 or raw:
                writer.writerow(["" if v is None else v for v in row])
            else:
                writer.writerow(
                    [format_value(header[i], v) for i, v in enumerate(row)]
                )
    return path


def render_text(table: Sequence[Sequence[Any]]) -> str:
    """Render a fixed-width text table for terminal output."""
    if not table:
        return ""

    header = list(table[0])
    rendered: List[List[str]] = [[str(h) for h in header]]
    for row in table[1:]:
        rendered.append([format_value(header[i], v) for i, v in enumerate(row)])

    widths = [
        max(len(rendered[r][c]) for r in range(len(rendered)))
        for c in range(len(header))
    ]

    lines = []
    for index, row in enumerate(rendered):
        lines.append("  ".join(cell.rjust(widths[i]) for i, cell in enumerate(row)))
        if index == 0:
            lines.append("  ".join("-" * w for w in widths))
    return "\n".join(lines)


def write_xlsx(
    table: Sequence[Sequence[Any]], path: str, sheet_title: str = "Comps"
) -> str:
    """Write a formatted Excel workbook. Requires openpyxl."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Excel export requires openpyxl, which is not installed.\n"
            "  pip install openpyxl\n"
            "Or export to CSV instead with --format csv."
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title

    header = list(table[0]) if table else []
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")

    for col_index, label in enumerate(header, start=1):
        cell = sheet.cell(row=1, column=col_index, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row_index, row in enumerate(table[1:], start=2):
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            label = header[col_index - 1] if col_index <= len(header) else ""
            if isinstance(value, (int, float)):
                if label in PERCENT_LABELS:
                    cell.number_format = "0.0%"
                elif label in MULTIPLE_LABELS:
                    cell.number_format = '0.0"x"'
                elif label in DECIMAL_LABELS:
                    cell.number_format = "#,##0.0"
                else:
                    cell.number_format = "#,##0"
            elif isinstance(value, str) and value in ("MEDIAN", "MEAN", "MIN", "MAX"):
                cell.font = Font(bold=True, italic=True)

    for col_index, label in enumerate(header, start=1):
        width = max(12, min(28, len(str(label)) + 4))
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    sheet.freeze_panes = "A2"
    workbook.save(path)
    return path


def write(
    table: Sequence[Sequence[Any]],
    path: str,
    fmt: Optional[str] = None,
    raw: bool = False,
) -> str:
    """Dispatch on format, inferring from the file extension when not given."""
    if fmt is None:
        fmt = "xlsx" if path.lower().endswith((".xlsx", ".xlsm")) else "csv"
    if fmt == "xlsx":
        return write_xlsx(table, path)
    if fmt == "csv":
        return write_csv(table, path, raw=raw)
    raise ValueError("Unsupported format {0!r}. Use 'csv' or 'xlsx'.".format(fmt))
